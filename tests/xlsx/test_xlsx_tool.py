"""The xlsx skill's helper, end to end.

Needs Python 3 with openpyxl and LibreOffice (`soffice`) on PATH. SKILL_DIR names the bundle (default
skills/xlsx next to this repository's tests). The fixtures are built at run time, so no binary file is committed.
"""

import json
import os
import subprocess
import sys
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl

TOOL = Path(os.environ.get("SKILL_DIR", Path(__file__).resolve().parents[2] / "skills" / "xlsx")) / "scripts" / "xlsx_tool.py"


def tool(*args: str) -> tuple[int, dict]:
    done = subprocess.run(
        [sys.executable, "-I", str(TOOL), *args],
        check=False,
        capture_output=True,
        text=True,
        timeout=240,
        stdin=subprocess.DEVNULL,
    )
    return done.returncode, json.loads(done.stdout)


class XlsxToolTest(unittest.TestCase):
    def setUp(self) -> None:
        self.dir = Path(tempfile.mkdtemp(prefix="xlsx-skill-"))

    def workbook(self, name: str = "in.xlsx") -> Path:
        """Numbers in A1:A3, a SUM and a division of it, and one deliberate #DIV/0!. openpyxl stores no cached values."""
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Data"
        for row, value in enumerate((10, 20, 30), start=1):
            sheet.cell(row=row, column=1, value=value)
        sheet["A4"] = "=SUM(A1:A3)"
        sheet["B4"] = "=A4/2"
        sheet["C1"] = "=1/0"
        path = self.dir / name
        book.save(path)
        return path

    def test_recalc_computes_formulas_written_without_cached_values(self) -> None:
        source, target = self.workbook(), self.dir / "out.xlsx"

        code, result = tool("recalc", str(source), str(target))

        self.assertEqual(code, 0, result)
        self.assertTrue(result["ok"])
        values = openpyxl.load_workbook(target, data_only=True)["Data"]
        self.assertEqual(values["A4"].value, 60)
        self.assertEqual(values["B4"].value, 30)
        self.assertEqual(values["C1"].value, "#DIV/0!")
        formulas = openpyxl.load_workbook(target, data_only=False)["Data"]
        self.assertEqual(formulas["A4"].value, "=SUM(A1:A3)", "formulas stay formulas")

    def test_check_reports_exactly_the_error_cells(self) -> None:
        source, target = self.workbook(), self.dir / "out.xlsx"
        code, result = tool("recalc", str(source), str(target))
        self.assertEqual(code, 0, result)

        code, result = tool("check", str(target))

        self.assertEqual(code, 1)
        self.assertFalse(result["ok"])
        self.assertEqual(
            result["errors"], [{"sheet": "Data", "cell": "C1", "error": "#DIV/0!", "formula": "=1/0"}], result
        )
        self.assertEqual(result["formulas_without_value"], [])
        self.assertEqual(result["formula_cells"], 3)

    def test_check_on_an_uncalculated_file_reports_every_formula(self) -> None:
        code, result = tool("check", str(self.workbook()))

        self.assertEqual(code, 1)
        self.assertEqual(result["errors"], [])
        self.assertEqual(sorted(entry["cell"] for entry in result["formulas_without_value"]), ["A4", "B4", "C1"])

    def test_check_is_clean_for_a_good_workbook(self) -> None:
        book = openpyxl.Workbook()
        book.active["A1"] = 2
        book.active["A2"] = "=A1*3"
        source, target = self.dir / "good.xlsx", self.dir / "good-out.xlsx"
        book.save(source)
        code, result = tool("recalc", str(source), str(target))
        self.assertEqual(code, 0, result)

        code, result = tool("check", str(target))

        self.assertEqual((code, result["ok"], result["error_count"]), (0, True, 0), result)

    def test_inspect_lists_sheets_dimensions_header_and_formula_count(self) -> None:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Sales"
        sheet.append(["Region", "Amount"])
        sheet.append(["North", 5])
        sheet.append(["Total", "=SUM(B2:B2)"])
        book.create_sheet("Empty")
        path = self.dir / "inspect.xlsx"
        book.save(path)

        code, result = tool("inspect", str(path))

        self.assertEqual(code, 0)
        sales, empty = result["sheets"]
        self.assertEqual(sales["name"], "Sales")
        self.assertEqual(sales["dimensions"], "A1:B3")
        self.assertEqual(sales["header"], ["Region", "Amount"])
        self.assertEqual(sales["formula_count"], 1)
        self.assertEqual(empty["header"], [])

    def test_recalc_reads_csv_and_writes_xlsx(self) -> None:
        source = self.dir / "table.csv"
        source.write_text("name,qty\nA,2\nB,3\n", encoding="utf-8")

        code, result = tool("recalc", str(source), str(self.dir / "table.xlsx"))

        self.assertEqual(code, 0, result)
        sheet = openpyxl.load_workbook(self.dir / "table.xlsx", data_only=True).active
        self.assertEqual(
            [[cell.value for cell in row] for row in sheet.iter_rows()], [["name", "qty"], ["A", 2], ["B", 3]]
        )

    def test_recalc_uses_a_private_profile_so_parallel_calls_succeed(self) -> None:
        def one(index: int) -> tuple[int, dict]:
            return tool("recalc", str(self.workbook(f"p{index}.xlsx")), str(self.dir / f"p{index}-out.xlsx"))

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(one, (1, 2)))

        for code, result in results:
            self.assertEqual(code, 0, result)

    def test_recalc_ends_with_a_json_error_when_the_timeout_hits(self) -> None:
        book = openpyxl.Workbook()
        sheet = book.active
        for row in range(1, 60001):
            sheet.append([row, f"=A{row}*2", f"=B{row}+A{row}"])
        big = self.dir / "big.xlsx"
        book.save(big)

        code, result = tool("recalc", str(big), str(self.dir / "big-out.xlsx"), "--timeout", "1")

        self.assertEqual(code, 1)
        self.assertIn("did not finish", result["error"])

    def test_a_missing_or_wrong_file_is_a_json_error(self) -> None:
        code, result = tool("recalc", str(self.dir / "nope.xlsx"), str(self.dir / "x.xlsx"))
        self.assertEqual((code, "does not exist" in result["error"]), (1, True))
        not_a_workbook = self.dir / "plain.xlsx"
        not_a_workbook.write_text("hello", encoding="utf-8")
        code, result = tool("check", str(not_a_workbook))
        self.assertEqual(code, 1)
        self.assertIn("recalc", result["error"])

    def test_output_opens_in_libreoffice(self) -> None:
        target = self.dir / "out.xlsx"
        tool("recalc", str(self.workbook()), str(target))
        profile = self.dir / "profile"

        command = ["soffice", f"-env:UserInstallation={profile.as_uri()}", "--headless", "--convert-to", "csv"]
        command += ["--outdir", str(self.dir), str(target)]
        for _round in range(2):  # a first start on a fresh profile may only create the profile
            done = subprocess.run(
                command, check=False, capture_output=True, text=True, timeout=170, stdin=subprocess.DEVNULL
            )
            if (self.dir / "out.csv").is_file():
                break

        self.assertTrue((self.dir / "out.csv").is_file(), done.stderr)
        self.assertIn("60", (self.dir / "out.csv").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
