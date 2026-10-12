#!/usr/bin/env python3
"""Calculate, check and inspect Excel workbooks.

Three subcommands, each printing one JSON object on stdout:

  recalc  IN OUT   calculate every formula and save OUT as .xlsx
  check   FILE     list error values and formulas that have no calculated value
  inspect FILE     sheets, sizes, header row and formula count

A failure prints {"error": "..."} and exits 1. `check` also exits 1 when it finds a problem
(its JSON says "ok": false), so `check FILE && echo fine` works in a shell.
"""

import argparse
import json
import os
import shutil
import signal
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ERROR_VALUES = (
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#SPILL!",
    "#CALC!",
    "#FIELD!",
    "#GETTING_DATA",
)
LIST_LIMIT = 200  # per list in a check result; the counts stay exact
HEADER_COLUMNS = 30

# Calc's "recalculation on file load" settings: 0 = always, 1 = never (the default), 2 = ask.
PROFILE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" \
xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse">\
<value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse">\
<value>0</value></prop></item>
</oor:items>
"""


class ToolError(Exception):
    """A failure the model can act on; its message goes to stdout as {"error": ...}."""


def recalc(source: str, target: str, timeout: int) -> dict:
    src = Path(source)
    if not src.is_file():
        raise ToolError(f"{source} does not exist")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise ToolError("LibreOffice is not available here")
    # A fresh profile per call: two calls never share a profile lock, and nothing of the user's is read.
    work = Path(tempfile.mkdtemp(prefix="xlsx-recalc-"))
    try:
        (work / "profile" / "user").mkdir(parents=True)
        (work / "profile" / "user" / "registrymodifications.xcu").write_text(PROFILE, encoding="utf-8")
        out_dir = work / "out"
        out_dir.mkdir()
        command = [
            soffice,
            f"-env:UserInstallation={(work / 'profile').as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            str(src.resolve()),
        ]
        produced = out_dir / (src.stem + ".xlsx")
        deadline = time.monotonic() + timeout
        log = ""
        # A first start on a fresh profile can end without converting anything, so one retry is allowed.
        for _round in range(2):
            process = subprocess.Popen(
                command,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                start_new_session=True,
            )
            try:
                log, _ = process.communicate(timeout=max(deadline - time.monotonic(), 1))
            except subprocess.TimeoutExpired:
                if hasattr(os, "killpg"):
                    os.killpg(process.pid, signal.SIGKILL)  # soffice starts a child of its own
                else:
                    process.kill()
                process.communicate()
                raise ToolError(f"LibreOffice did not finish within {timeout} s; try a smaller workbook or --timeout")
            if process.returncode == 0 and produced.is_file():
                break
        if not produced.is_file():
            raise ToolError(f"LibreOffice could not convert {source}: {log.strip()[-300:]}")
        destination = Path(target)
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), destination)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return {"ok": True, "output": str(destination)}


def load(path, data_only: bool):
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=data_only)
    except Exception as error:  # noqa: BLE001 - a wrong format or a damaged file: say so, do not dump a traceback
        raise ToolError(f"{path} is not an .xlsx workbook openpyxl can read ({type(error).__name__}); use recalc first")


def check(path: str) -> dict:
    formulas = load(path, data_only=False)
    values = load(path, data_only=True)
    errors: list[dict] = []
    without_value: list[dict] = []
    error_count = without_count = formula_count = 0
    for sheet in formulas.worksheets:
        calculated = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                is_formula = cell.data_type == "f"
                formula_count += is_formula
                result = calculated[cell.coordinate]
                value = result.value
                # A text cell that merely reads "#N/A" is data, not an error: only formulas and real error cells count.
                if isinstance(value, str) and value in ERROR_VALUES and (is_formula or result.data_type == "e"):
                    error_count += 1
                    if len(errors) < LIST_LIMIT:
                        entry = {"sheet": sheet.title, "cell": cell.coordinate, "error": value}
                        if is_formula:
                            entry["formula"] = str(cell.value)
                        errors.append(entry)
                elif is_formula and value is None:
                    without_count += 1
                    if len(without_value) < LIST_LIMIT:
                        without_value.append(
                            {"sheet": sheet.title, "cell": cell.coordinate, "formula": str(cell.value)}
                        )
    return {
        "ok": error_count == 0 and without_count == 0,
        "formula_cells": formula_count,
        "error_count": error_count,
        "errors": errors,
        "without_value_count": without_count,
        "formulas_without_value": without_value,
        "truncated": error_count > len(errors) or without_count > len(without_value),
    }


def inspect(path: str) -> dict:
    workbook = load(path, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        header: list = []
        dimensions = sheet.dimensions  # read first: iterating past the used range would widen it
        columns = sheet.max_column
        for row in sheet.iter_rows(max_col=min(columns, HEADER_COLUMNS), values_only=True):
            if any(value is not None for value in row):
                header = [None if value is None else str(value) for value in row]
                break
        formula_count = sum(1 for row in sheet.iter_rows() for cell in row if cell.data_type == "f")
        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimensions": dimensions,
                "rows": sheet.max_row,
                "columns": columns,
                "header": header,
                "formula_count": formula_count,
            }
        )
    return {"sheets": sheets}


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    one = commands.add_parser("recalc", help="calculate every formula; reads .xlsx, .xls, .csv or .ods, writes .xlsx")
    one.add_argument("input")
    one.add_argument("output")
    one.add_argument("--timeout", type=int, default=120, help="seconds before LibreOffice is stopped (default 120)")
    commands.add_parser("check", help="error values and formulas without a value").add_argument("file")
    commands.add_parser("inspect", help="sheets, sizes, header row, formula count").add_argument("file")
    args = parser.parse_args(argv)
    try:
        if args.command == "recalc":
            result = recalc(args.input, args.output, args.timeout)
        elif args.command == "check":
            result = check(args.file)
        else:
            result = inspect(args.file)
    except ToolError as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result.get("ok", True) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
